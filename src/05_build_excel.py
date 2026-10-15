import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule

Path("excel").mkdir(exist_ok=True)
cls = pd.read_csv("data/processed/sku_classification.csv")
fc = pd.read_csv("data/processed/forecast_summary.csv")
risk = pd.read_csv("data/processed/stockout_risk.csv")

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")

def style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    ws.freeze_panes = "A2"

def write_df(ws, d):
    ws.append(list(d.columns))
    for row in d.itertuples(index=False):
        ws.append(list(row))
    style_header(ws)

wb = Workbook()

# ---- Dashboard (filled after other sheets exist) ----
ds = wb.active
ds.title = "Dashboard"

# ---- Assumptions: change these and the Risk sheet recalculates ----
asm = wb.create_sheet("Assumptions")
asm.append(["Parameter", "Value", "Meaning"])
asm.append(["Lead time (days)", 3, "Days from placing an order to receiving stock"])
asm.append(["Service level z", 1.65, "1.65 = 95% chance of not running out"])
asm.append(["Review buffer (days)", 7, "Extra days of demand to order for"])
style_header(asm)
asm.column_dimensions["A"].width = 24
asm.column_dimensions["C"].width = 50

# ---- Classification ----
write_df(wb.create_sheet("Classification"), cls)

# ---- Forecast ----
write_df(wb.create_sheet("Forecast"), fc)

# ---- Risk sheet with live formulas ----
rs = wb.create_sheet("Risk")
rs.append(["sku", "current_stock", "forecast_daily", "std_daily", "days_cover",
           "safety_stock", "reorder_point", "risk", "order_qty"])
for i, row in enumerate(risk.itertuples(index=False), start=2):
    rs.append([row.sku, row.current_stock, row.forecast_daily, row.std_daily,
               f"=IF(C{i}>0,B{i}/C{i},999)",
               f"=Assumptions!$B$3*D{i}*SQRT(Assumptions!$B$2)",
               f"=C{i}*Assumptions!$B$2+F{i}",
               f'=IF(E{i}<Assumptions!$B$2,"High",IF(B{i}<G{i},"Medium","Low"))',
               f"=MAX(0,ROUNDUP(G{i}+C{i}*Assumptions!$B$4-B{i},0))"])
style_header(rs)
last = len(risk) + 1
rs.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=['"High"'],
                              fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))
rs.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=['"Medium"'],
                              fill=PatternFill("solid", bgColor="FFE699", fgColor="FFE699")))
rs.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=['"Low"'],
                              fill=PatternFill("solid", bgColor="C6E0B4", fgColor="C6E0B4")))

# ---- Dashboard content ----
ds["A1"] = "Inventory & Demand Planning Dashboard"
ds["A1"].font = Font(bold=True, size=16)
ds.append([])
ds["A3"], ds["B3"] = "Stock-out risk", "SKUs"
for i, lvl in enumerate(["High", "Medium", "Low"], start=4):
    ds[f"A{i}"] = lvl
    ds[f"B{i}"] = f'=COUNTIF(Risk!$H:$H,A{i})'
ds["A8"], ds["B8"] = "Velocity", "SKUs"
for i, v in enumerate(["Fast", "Medium", "Slow"], start=9):
    ds[f"A{i}"] = v
    ds[f"B{i}"] = f'=COUNTIF(Classification!$L:$L,A{i})'
for cell in ("A3", "B3", "A8", "B8"):
    ds[cell].font = Font(bold=True)

c1 = BarChart(); c1.title = "SKUs by stock-out risk"
c1.add_data(Reference(ds, min_col=2, min_row=3, max_row=6), titles_from_data=True)
c1.set_categories(Reference(ds, min_col=1, min_row=4, max_row=6))
ds.add_chart(c1, "D3")

c2 = BarChart(); c2.title = "SKUs by velocity"
c2.add_data(Reference(ds, min_col=2, min_row=8, max_row=11), titles_from_data=True)
c2.set_categories(Reference(ds, min_col=1, min_row=9, max_row=11))
ds.add_chart(c2, "D20")

wb.save("excel/inventory_planning_model.xlsx")
print("Saved excel/inventory_planning_model.xlsx")