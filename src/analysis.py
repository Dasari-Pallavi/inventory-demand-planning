"""Inventory & Demand Planning Analysis
Run from the project root:  python src/analysis.py
"""
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.drawing.image import Image as XLImage

# ---------------- CONFIG (assumptions, change freely) ----------------
RAW = Path("data/raw/retail_store_inventory.csv")
OUT = Path("outputs"); CH = OUT / "charts"
LEAD_TIME = 7          # supplier lead time in days
Z = 1.65               # service level z-score (95%)
OVERSTOCK_DAYS = 60    # more than this many days of cover = overstock
HORIZON = 14           # forecast days
HOLDOUT = 28           # backtest days
WINDOWS = [7, 14, 30]  # moving-average windows to compare
CH.mkdir(parents=True, exist_ok=True)
Path("data/processed").mkdir(parents=True, exist_ok=True)


def save(name):
    plt.tight_layout()
    plt.savefig(CH / name, dpi=150)
    plt.close()


def wape(actual, pred):
    return np.abs(actual - pred).sum() / max(actual.sum(), 1e-9)


# ---------------- 1. LOAD & CLEAN ----------------
df = pd.read_csv(RAW)
df.columns = df.columns.str.strip().str.lower().str.replace(r"[ /]+", "_", regex=True)
df["date"] = pd.to_datetime(df["date"])
print("Shape:", df.shape)
print(df.isna().sum())
print(df.describe())

df = df.drop_duplicates().dropna(subset=["date", "product_id", "units_sold"])
df = df[(df["units_sold"] >= 0) & (df["inventory_level"] >= 0)]
df["revenue"] = df["units_sold"] * df["price"]
df.to_csv("data/processed/clean_sales.csv", index=False)

# ---------------- 2. DAILY SKU-LEVEL DATA ----------------
daily = (df.groupby(["date", "product_id"], as_index=False)
           .agg(units_sold=("units_sold", "sum"), inventory_level=("inventory_level", "sum")))
demand = (daily.pivot(index="date", columns="product_id", values="units_sold")
               .asfreq("D").fillna(0))

# ---------------- 3. SKU SUMMARY: VELOCITY + ABC ----------------
sku = (df.groupby("product_id")
         .agg(category=("category", "first"),
              total_units=("units_sold", "sum"),
              revenue=("revenue", "sum"))
         .join(daily.groupby("product_id")["inventory_level"].mean().rename("avg_inventory")))
sku["avg_daily_demand"] = demand.mean()
sku["std_daily_demand"] = demand.std()
sku["turnover"] = sku["total_units"] / sku["avg_inventory"]

q75, q25 = sku["avg_daily_demand"].quantile([.75, .25])
sku["velocity"] = np.select([sku["avg_daily_demand"] >= q75, sku["avg_daily_demand"] <= q25],
                            ["Fast", "Slow"], "Medium")

sku = sku.sort_values("revenue", ascending=False)
sku["cum_share"] = sku["revenue"].cumsum() / sku["revenue"].sum()
sku["abc"] = np.select([sku["cum_share"] <= .80, sku["cum_share"] <= .95], ["A", "B"], "C")

# ---------------- 4. FORECAST (MOVING AVERAGE) + BACKTEST ----------------
def backtest(series, w):
    pred = series.shift(1).rolling(w).mean()      # one-step-ahead rolling forecast
    return wape(series.iloc[-HOLDOUT:].values, pred.iloc[-HOLDOUT:].values)

fc_rows, fc_long = [], []
last_date = demand.index.max()
future = pd.date_range(last_date + pd.Timedelta(days=1), periods=HORIZON)
for p in demand.columns:
    s = demand[p]
    scores = {w: backtest(s, w) for w in WINDOWS}
    best = min(scores, key=scores.get)
    fc = s.tail(best).mean()
    fc_rows.append({"product_id": p, "best_window": best, "wape": scores[best],
                    "forecast_daily": fc, "forecast_14d_units": fc * HORIZON})
    fc_long += [{"product_id": p, "date": d, "forecast_units": round(fc, 2)} for d in future]
fc_df = pd.DataFrame(fc_rows).set_index("product_id")
fc_long = pd.DataFrame(fc_long)

# ---------------- 5. STOCK-OUT RISK ----------------
latest = daily[daily["date"] == daily["date"].max()].set_index("product_id")["inventory_level"]
report = sku.join(fc_df)
report["current_stock"] = latest
report["days_of_cover"] = report["current_stock"] / report["forecast_daily"].clip(lower=0.01)
report["safety_stock"] = Z * report["std_daily_demand"] * np.sqrt(LEAD_TIME)
report["reorder_point"] = report["forecast_daily"] * LEAD_TIME + report["safety_stock"]
report["risk"] = np.select(
    [report["days_of_cover"] < LEAD_TIME,
     report["current_stock"] < report["reorder_point"],
     report["days_of_cover"] > OVERSTOCK_DAYS],
    ["High", "Medium", "Overstock"], "Healthy")
report["suggested_order_qty"] = np.where(
    report["risk"].isin(["High", "Medium"]),
    (report["reorder_point"] + report["forecast_daily"] * HORIZON - report["current_stock"]).clip(lower=0).round(),
    0)
report = report.reset_index()

report.to_csv(OUT / "stockout_risk.csv", index=False)
fc_long.to_csv(OUT / "forecast_14d.csv", index=False)
print("\nRisk summary:\n", report["risk"].value_counts())
print("Average WAPE:", round(report["wape"].mean(), 3))

# ---------------- 6. CHARTS (EDA + results) ----------------
tot = daily.groupby("date")["units_sold"].sum()
plt.figure(figsize=(11, 4))
tot.plot(alpha=.4, label="Daily units")
tot.rolling(30).mean().plot(label="30-day average", linewidth=2)
plt.title("Total Units Sold Over Time"); plt.ylabel("Units"); plt.legend()
save("01_sales_trend.png")

plt.figure(figsize=(7, 4))
df.groupby("category")["units_sold"].sum().sort_values().plot.barh(color="#4C78A8")
plt.title("Units Sold by Category"); save("02_category_sales.png")

plt.figure(figsize=(8, 4))
df.groupby(df["date"].dt.month)["units_sold"].sum().plot(marker="o")
plt.title("Monthly Seasonality (Units Sold)"); plt.xlabel("Month"); plt.xticks(range(1, 13))
save("03_monthly_seasonality.png")

ranked = sku.sort_values("total_units")
top_bottom = pd.concat([ranked.head(10), ranked.tail(10)])
plt.figure(figsize=(8, 6))
plt.barh(top_bottom.index, top_bottom["total_units"],
         color=["#E45756"] * min(10, len(ranked)) + ["#54A24B"] * min(10, len(ranked)))
plt.title("Slowest (red) vs Fastest (green) SKUs - Total Units"); save("04_fast_vs_slow.png")

fig, ax = plt.subplots(figsize=(10, 4))
ax.bar(sku.index, sku["revenue"], color="#4C78A8")
ax.tick_params(axis="x", rotation=90)
ax2 = ax.twinx(); ax2.plot(sku.index, sku["cum_share"] * 100, color="#F58518", marker="o")
ax2.axhline(80, ls="--", color="grey"); ax2.set_ylabel("Cumulative %")
ax.set_title("ABC Pareto: Revenue by SKU"); save("05_pareto_abc.png")

top_sku = sku["total_units"].idxmax()
s = demand[top_sku]
best_w = int(fc_df.loc[top_sku, "best_window"])
plt.figure(figsize=(11, 4))
s.tail(90).plot(label="Actual")
s.shift(1).rolling(best_w).mean().tail(90).plot(label=f"{best_w}-day MA (backtest)")
plt.plot(future, [fc_df.loc[top_sku, "forecast_daily"]] * HORIZON, "r--", label="Forecast")
plt.title(f"Forecast vs Actual - {top_sku}"); plt.legend(); save("06_forecast_vs_actual.png")

colors = {"High": "#E45756", "Medium": "#F58518", "Healthy": "#54A24B", "Overstock": "#4C78A8"}
rc = report["risk"].value_counts()
plt.figure(figsize=(6, 4))
plt.bar(rc.index, rc.values, color=[colors[k] for k in rc.index])
plt.title("SKUs by Stock Risk Status"); save("07_risk_status.png")

num = df.select_dtypes("number").corr()
plt.figure(figsize=(7, 6))
plt.imshow(num, cmap="coolwarm", vmin=-1, vmax=1); plt.colorbar()
plt.xticks(range(len(num)), num.columns, rotation=90); plt.yticks(range(len(num)), num.columns)
plt.title("Correlation Heatmap"); save("08_correlation.png")

# ---------------- 7. EXCEL MODEL (live formulas) ----------------
cols = ["product_id", "category", "velocity", "abc", "total_units", "revenue",
        "avg_daily_demand", "turnover", "best_window", "wape", "risk", "suggested_order_qty"]
fills = {"High": "F8CBAD", "Medium": "FFE699", "Healthy": "C6E0B4", "Overstock": "BDD7EE"}

with pd.ExcelWriter(OUT / "inventory_model.xlsx", engine="openpyxl") as w:
    pd.DataFrame({"Parameter": ["Lead time (days)", "Z (service level)", "Overstock threshold (days)"],
                  "Value": [LEAD_TIME, Z, OVERSTOCK_DAYS]}).to_excel(w, sheet_name="Params", index=False)
    report[cols].round(2).to_excel(w, sheet_name="SKU_Summary", index=False)
    fc_long.to_excel(w, sheet_name="Forecast_14d", index=False)

    ws = w.book.create_sheet("Stockout_Model")
    ws.append(["product_id", "category", "current_stock", "forecast_daily", "std_daily",
               "days_of_cover", "safety_stock", "reorder_point", "risk"])
    for i, r in enumerate(report.itertuples(), start=2):
        ws.append([r.product_id, r.category, int(r.current_stock), round(r.forecast_daily, 2),
                   round(r.std_daily_demand, 2),
                   f"=IF(D{i}=0,999,C{i}/D{i})",
                   f"=Params!$B$3*E{i}*SQRT(Params!$B$2)",
                   f"=D{i}*Params!$B$2+G{i}",
                   f'=IF(F{i}<Params!$B$2,"High",IF(C{i}<H{i},"Medium",IF(F{i}>Params!$B$4,"Overstock","Healthy")))'])
    rng_ = f"I2:I{len(report) + 1}"
    for label, color in fills.items():
        ws.conditional_formatting.add(rng_, CellIsRule(operator="equal", formula=[f'"{label}"'],
                                      fill=PatternFill(start_color=color, end_color=color, fill_type="solid")))

    dash = w.book.create_sheet("Dashboard")
    dash["A1"] = "Inventory & Demand Planning Dashboard"
    for img, cell in [("07_risk_status.png", "A3"), ("05_pareto_abc.png", "J3"),
                      ("06_forecast_vs_actual.png", "A25"), ("04_fast_vs_slow.png", "J25")]:
        pic = XLImage(str(CH / img)); pic.width, pic.height = pic.width * .35, pic.height * .35
        dash.add_image(pic, cell)

print("Done. Check the outputs/ folder.")